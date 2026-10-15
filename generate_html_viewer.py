import openpyxl
import json
import html
from pathlib import Path

wb_path = Path("c:/Users/pankh/OneDrive/Desktop/Govt/NESFIC_D36_Government_AI_Dataset_Master.xlsx")
wb = openpyxl.load_workbook(wb_path, data_only=True)

sheets_data = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = []
    for r in range(1, ws.max_row + 1):
        row_vals = []
        is_empty = True
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is not None and str(val).strip() != "":
                is_empty = False
            row_vals.append("" if val is None else str(val))
        if not is_empty:
            rows.append(row_vals)
    sheets_data[sheet_name] = rows

html_template = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>NESFIC-D-36 Master Government AI Dataset</title>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
<style>
  :root {
    --primary: #1B365D;
    --primary-light: #2A4D7C;
    --accent: #D9A74A;
    --bg: #F8FAFC;
    --card-bg: #FFFFFF;
    --border: #E2E8F0;
    --text: #1E293B;
    --text-muted: #64748B;
    --green-bg: #E2EFDA;
    --green-text: #276A3C;
    --amber-bg: #FFF2CC;
    --amber-text: #8C6B00;
    --red-bg: #FCE4D6;
    --red-text: #C00000;
  }
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body {
    font-family: 'Inter', -apple-system, sans-serif;
    background: var(--bg);
    color: var(--text);
    display: flex;
    height: 100vh;
    overflow: hidden;
  }
  #sidebar {
    width: 320px;
    background: #0F172A;
    color: #F8FAFC;
    display: flex;
    flex-direction: column;
    border-right: 1px solid #1E293B;
    flex-shrink: 0;
  }
  .sidebar-header {
    padding: 20px;
    border-bottom: 1px solid #1E293B;
    background: #090D16;
  }
  .sidebar-header h1 {
    font-size: 15px;
    font-weight: 700;
    color: #FFFFFF;
    letter-spacing: -0.3px;
    margin-bottom: 4px;
  }
  .sidebar-header p {
    font-size: 11px;
    color: #94A3B8;
  }
  .badge {
    display: inline-block;
    padding: 2px 8px;
    border-radius: 4px;
    font-size: 10px;
    font-weight: 600;
    margin-top: 8px;
    background: rgba(217, 167, 74, 0.2);
    color: #F6D365;
    border: 1px solid rgba(217, 167, 74, 0.4);
  }
  .sheet-list {
    flex: 1;
    overflow-y: auto;
    padding: 12px 8px;
  }
  .sheet-item {
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 9px 12px;
    border-radius: 6px;
    font-size: 12.5px;
    font-weight: 500;
    color: #CBD5E1;
    cursor: pointer;
    margin-bottom: 3px;
    transition: all 0.15s ease;
  }
  .sheet-item:hover {
    background: rgba(255, 255, 255, 0.06);
    color: #FFFFFF;
  }
  .sheet-item.active {
    background: #1B365D;
    color: #FFFFFF;
    font-weight: 600;
    box-shadow: 0 2px 6px rgba(0,0,0,0.3);
  }
  .sheet-count {
    font-size: 10px;
    padding: 2px 6px;
    border-radius: 10px;
    background: rgba(255, 255, 255, 0.1);
    color: #94A3B8;
    font-family: 'JetBrains Mono', monospace;
  }
  #main {
    flex: 1;
    display: flex;
    flex-direction: column;
    overflow: hidden;
  }
  .top-bar {
    height: 64px;
    background: var(--card-bg);
    border-bottom: 1px solid var(--border);
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 0 24px;
    flex-shrink: 0;
  }
  .top-bar-title h2 {
    font-size: 17px;
    font-weight: 700;
    color: var(--primary);
  }
  .top-bar-title span {
    font-size: 12px;
    color: var(--text-muted);
    margin-left: 10px;
    font-weight: 500;
  }
  .search-box {
    display: flex;
    align-items: center;
    gap: 8px;
  }
  .search-box input {
    padding: 7px 14px;
    border: 1px solid var(--border);
    border-radius: 6px;
    font-size: 12.5px;
    width: 260px;
    font-family: inherit;
    outline: none;
    transition: border-color 0.2s;
  }
  .search-box input:focus {
    border-color: var(--primary);
  }
  .content-area {
    flex: 1;
    overflow: auto;
    padding: 20px 24px;
  }
  .table-wrapper {
    background: var(--card-bg);
    border: 1px solid var(--border);
    border-radius: 8px;
    overflow: auto;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04);
  }
  table {
    width: 100%;
    border-collapse: collapse;
    font-size: 12px;
    white-space: nowrap;
  }
  th {
    background: #1B365D;
    color: #FFFFFF;
    font-weight: 600;
    padding: 10px 14px;
    text-align: left;
    position: sticky;
    top: 0;
    z-index: 10;
    border-right: 1px solid rgba(255, 255, 255, 0.1);
  }
  td {
    padding: 9px 14px;
    border-bottom: 1px solid var(--border);
    border-right: 1px solid var(--border);
    color: var(--text);
    vertical-align: top;
    white-space: normal;
    max-width: 420px;
    line-height: 1.45;
  }
  tr:nth-child(even) td {
    background: #F8FAFC;
  }
  tr:hover td {
    background: #F1F5F9;
  }
  /* Status chips */
  .tag {
    display: inline-block;
    padding: 2px 7px;
    border-radius: 4px;
    font-size: 10px;
    font-weight: 700;
    text-transform: uppercase;
  }
  .tag-green { background: var(--green-bg); color: var(--green-text); }
  .tag-amber { background: var(--amber-bg); color: var(--amber-text); }
  .tag-red { background: var(--red-bg); color: var(--red-text); }
  .tag-code {
    font-family: 'JetBrains Mono', monospace;
    font-size: 11px;
    background: #EDF2F7;
    padding: 1px 5px;
    border-radius: 3px;
    color: #2D3748;
  }
</style>
</head>
<body>

<div id="sidebar">
  <div class="sidebar-header">
    <h1>NESFIC-D-36 MASTER DATASET</h1>
    <p>Trusted Government Knowledge Assistant</p>
    <div class="badge">19 SHEETS • 790 PAGES AUDITED</div>
  </div>
  <div class="sheet-list" id="sheetTabs"></div>
</div>

<div id="main">
  <div class="top-bar">
    <div class="top-bar-title">
      <h2 id="activeSheetTitle">README & Dataset Guide</h2>
      <span id="activeSheetMeta">Sheet 1 of 19</span>
    </div>
    <div class="search-box">
      <input type="text" id="filterInput" placeholder="Filter rows in active sheet..." onkeyup="filterTable()">
    </div>
  </div>
  <div class="content-area">
    <div class="table-wrapper">
      <table id="dataTable">
        <thead id="tableHead"></thead>
        <tbody id="tableBody"></tbody>
      </table>
    </div>
  </div>
</div>

<script>
const DATA = """ + json.dumps(sheets_data) + """;

const sheetNames = Object.keys(DATA);
let currentSheet = sheetNames[0];

function renderSidebar() {
  const container = document.getElementById("sheetTabs");
  container.innerHTML = "";
  sheetNames.forEach((name, idx) => {
    const item = document.createElement("div");
    item.className = "sheet-item" + (name === currentSheet ? " active" : "");
    const rowCount = DATA[name].length;
    item.innerHTML = `<span>${idx + 1}. ${name}</span><span class="sheet-count">${rowCount}</span>`;
    item.onclick = () => selectSheet(name, idx);
    container.appendChild(item);
  });
}

function formatCellValue(val) {
  if (!val) return "";
  const upper = val.trim().toUpperCase();
  if (["VERIFIED", "ACTIVE", "SUPPORTED", "PASSED", "APPROVED", "VALID"].includes(upper)) {
    return `<span class="tag tag-green">${val}</span>`;
  }
  if (["PENDING_VERIFICATION", "PENDING_APPROVAL", "VERIFICATION_REQUIRED", "UNCONFIRMED_CURRENT_VALIDITY", "PROVISIONAL", "SUSPECT_ORIGIN_THIRD_PARTY_TOOL"].includes(upper)) {
    return `<span class="tag tag-amber">${val}</span>`;
  }
  if (["SUPERSEDED", "REVOKED", "UNSUPPORTED", "CONTRADICTED", "CANNOT_GENERATE", "FAILED", "REJECTED_QUARANTINED", "EXCLUDED_PII"].includes(upper)) {
    return `<span class="tag tag-red">${val}</span>`;
  }
  if (val.startsWith("DOC-") || val.startsWith("SRC-") || val.startsWith("VER-") || val.startsWith("EVD-") || val.startsWith("CHK-") || val.startsWith("RET-") || val.startsWith("EVAL-") || val.startsWith("ABS-") || val.startsWith("ISSUE-")) {
    return `<span class="tag-code">${val}</span>`;
  }
  return val;
}

function selectSheet(name, idx) {
  currentSheet = name;
  document.getElementById("activeSheetTitle").innerText = name;
  document.getElementById("activeSheetMeta").innerText = `Sheet ${idx + 1} of ${sheetNames.length} • ${DATA[name].length} rows`;
  document.getElementById("filterInput").value = "";
  renderSidebar();
  renderTable(DATA[name]);
}

function renderTable(rows) {
  const thead = document.getElementById("tableHead");
  const tbody = document.getElementById("tableBody");
  thead.innerHTML = "";
  tbody.innerHTML = "";
  
  if (!rows || rows.length === 0) return;
  
  // Header row
  const hTr = document.createElement("tr");
  rows[0].forEach(cell => {
    const th = document.createElement("th");
    th.innerText = cell;
    hTr.appendChild(th);
  });
  thead.appendChild(hTr);
  
  // Body rows
  for (let r = 1; r < rows.length; r++) {
    const tr = document.createElement("tr");
    rows[r].forEach(cell => {
      const td = document.createElement("td");
      td.innerHTML = formatCellValue(cell);
      tr.appendChild(td);
    });
    tbody.appendChild(tr);
  }
}

function filterTable() {
  const query = document.getElementById("filterInput").value.toLowerCase();
  const rows = document.querySelectorAll("#tableBody tr");
  rows.forEach(tr => {
    const text = tr.innerText.toLowerCase();
    tr.style.display = text.includes(query) ? "" : "none";
  });
}

// Initial render
renderSidebar();
selectSheet(currentSheet, 0);
</script>
</body>
</html>
"""

viewer_path = Path("c:/Users/pankh/OneDrive/Desktop/Govt/view_dataset.html")
with open(viewer_path, "w", encoding="utf-8") as f:
    f.write(html_template)

print(f"Interactive HTML dataset viewer generated at: {viewer_path} ({viewer_path.stat().st_size} bytes)")
