import openpyxl, os, csv

wb = openpyxl.load_workbook('NESFIC_D36_Government_AI_Dataset_Master.xlsx', data_only=True)
out_dir = 'dataset_csv_sheets'
os.makedirs(out_dir, exist_ok=True)

exported = []
for i, sheetname in enumerate(wb.sheetnames, 1):
    ws = wb[sheetname]
    clean_name = sheetname.replace(" ", "_").replace("&", "and")
    safe_name = f"{i:02d}_{clean_name}.csv"
    csv_path = os.path.join(out_dir, safe_name)
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                writer.writerow([cell if cell is not None else '' for cell in row])
    exported.append((safe_name, ws.max_row, ws.max_column))
    print(f"Exported: {safe_name} ({ws.max_row} rows, {ws.max_column} cols)")

print(f"\nSuccessfully exported all {len(exported)} sheets to {out_dir}/")
