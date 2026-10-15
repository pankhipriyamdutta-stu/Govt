import openpyxl
from pathlib import Path

wb_path = Path("c:/Users/pankh/OneDrive/Desktop/Govt/NESFIC_D36_Government_AI_Dataset_Master.xlsx")
wb = openpyxl.load_workbook(wb_path)

print(f"Workbook loaded successfully: {wb_path.name}")
print(f"Total Sheets: {len(wb.sheetnames)}")
print("="*60)
for idx, name in enumerate(wb.sheetnames, 1):
    ws = wb[name]
    max_r = ws.max_row
    max_c = ws.max_column
    headers = [str(ws.cell(1, col).value) for col in range(1, min(6, max_c + 1))]
    print(f"{idx:2d}. Sheet: '{name:<32}' | Rows: {max_r:3d} | Cols: {max_c:2d} | Sample Cols: {', '.join(headers[:3])}")
print("="*60)
