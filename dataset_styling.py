"""
Master Dataset Builder for NESFIC-D-36
Trusted Government Knowledge, Rules & Document Assistant
Challenge: NESFIC 2026, Challenge Stream 1
"""

import sys
import os
import json
import hashlib
import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

# Palettes & Styling Constants
NAVY_HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1B365D")
SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="4A5568")
SECTION_FONT = Font(name="Segoe UI", size=11, bold=True, color="1B365D")
BODY_FONT = Font(name="Segoe UI", size=9.5, color="1A202C")
BOLD_BODY_FONT = Font(name="Segoe UI", size=9.5, bold=True, color="1A202C")
CODE_FONT = Font(name="Consolas", size=9, color="2D3748")

# Status Fills
STATUS_FILLS = {
    "VERIFIED": (PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="276A3C")),
    "ACTIVE": (PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="276A3C")),
    "PASSED": (PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="276A3C")),
    "SUPPORTED": (PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="276A3C")),
    "APPROVED": (PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="276A3C")),
    "VALID": (PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="276A3C")),
    
    "PENDING_VERIFICATION": (PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="8C6B00")),
    "PENDING_APPROVAL": (PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="8C6B00")),
    "VERIFICATION_REQUIRED": (PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="8C6B00")),
    "OCR_REVIEW_REQUIRED": (PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="8C6B00")),
    "PARTIALLY_SUPPORTED": (PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="8C6B00")),
    "INCOMPLETE": (PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="8C6B00")),
    "HUMAN_REVIEW_MANDATORY": (PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="8C6B00")),
    
    "SUPERSEDED": (PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="C00000")),
    "REVOKED": (PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="C00000")),
    "UNSUPPORTED": (PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="C00000")),
    "CONTRADICTED": (PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="C00000")),
    "CANNOT_GENERATE": (PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="C00000")),
    "FAILED": (PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="C00000")),
    "CONFLICTING": (PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="C00000")),
    
    "ARCHIVED": (PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"), Font(name="Segoe UI", size=9.5, color="595959")),
    "NOT_APPLICABLE": (PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"), Font(name="Segoe UI", size=9.5, color="595959")),
    "N/A": (PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"), Font(name="Segoe UI", size=9.5, color="595959")),
    "NOT_ASSESSED": (PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"), Font(name="Segoe UI", size=9.5, color="595959")),
}

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

def style_table(ws, start_row=1):
    """Applies clean borders, alignments, and widths to a sheet."""
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    
    # Auto-filter on headers
    max_col = ws.max_column
    max_row = ws.max_row
    if max_col > 0 and max_row >= start_row:
        start_letter = get_column_letter(1)
        end_letter = get_column_letter(max_col)
        ws.auto_filter.ref = f"{start_letter}{start_row}:{end_letter}{max_row}"
        
    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(start_row, max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            val_str = str(val) if val is not None else ""
            
            # Format header
            if row_idx == start_row:
                cell.fill = NAVY_HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.row_dimensions[row_idx].height = 28
                max_len = max(max_len, len(val_str))
            else:
                # Zebra fill
                if cell.fill.fill_type is None:
                    cell.fill = ZEBRA_FILL if row_idx % 2 == 0 else WHITE_FILL
                cell.border = THIN_BORDER
                
                # Check status highlighting
                val_upper = val_str.strip().upper()
                if val_upper in STATUS_FILLS:
                    fill, font = STATUS_FILLS[val_upper]
                    cell.fill = fill
                    cell.font = font
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                else:
                    if not cell.font or cell.font.name != "Segoe UI":
                        cell.font = BODY_FONT
                    # Alignment heuristics
                    if val_str.startswith(("DOC-", "SRC-", "VER-", "PG-", "CHK-", "EVD-", "CLM-", "REL-", "RET-", "EVAL-", "ABS-", "INT-", "ROL-", "SEC-", "LOG-", "AUD-", "ISSUE-")):
                        cell.alignment = Alignment(horizontal="center", vertical="top")
                    elif len(val_str) == 10 and val_str.count("-") == 2: # ISO date
                        cell.alignment = Alignment(horizontal="center", vertical="top")
                    elif len(val_str) > 40:
                        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="top")
                        
                max_len = max(max_len, min(len(val_str), 50))
                
        # Set column width
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        if ws.column_dimensions[col_letter].width > 60:
            ws.column_dimensions[col_letter].width = 60

print("Helper styling loaded.")
