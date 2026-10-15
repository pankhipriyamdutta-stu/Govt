import openpyxl, os

wb = openpyxl.load_workbook('NESFIC_D36_Government_AI_Dataset_Master.xlsx', data_only=True)
md_path = 'DATASET_WORKBOOK_VIEWER.md'

with open(md_path, 'w', encoding='utf-8') as f:
    f.write("# NESFIC-D-36: Master Government AI Dataset Workbook Viewer\n\n")
    f.write("> **Direct File Links:**\n")
    f.write("> - Master Excel File: [NESFIC_D36_Government_AI_Dataset_Master.xlsx](file:///c:/Users/pankh/OneDrive/Desktop/Govt/NESFIC_D36_Government_AI_Dataset_Master.xlsx)\n")
    f.write("> - Interactive Browser Viewer: [view_dataset.html](file:///c:/Users/pankh/OneDrive/Desktop/Govt/view_dataset.html)\n")
    f.write("> - Individual Sheet CSVs: [dataset_csv_sheets/](file:///c:/Users/pankh/OneDrive/Desktop/Govt/dataset_csv_sheets/)\n\n")
    
    f.write("## Table of Contents (All 19 Sheets)\n\n")
    for i, name in enumerate(wb.sheetnames, 1):
        anchor = name.lower().replace(" ", "-").replace("&", "").replace("---", "-")
        f.write(f"{i}. [{name}](#sheet-{i:02d}-{anchor})\n")
    f.write("\n---\n\n")

    for i, name in enumerate(wb.sheetnames, 1):
        ws = wb[name]
        anchor = name.lower().replace(" ", "-").replace("&", "").replace("---", "-")
        f.write(f"## Sheet {i:02d}: {name}\n\n")
        f.write(f"*Total Rows: {ws.max_row} | Columns: {ws.max_column}*\n\n")
        
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            f.write("*Empty sheet*\n\n---\n\n")
            continue
            
        headers = [str(c) if c is not None else '' for c in rows[0]]
        f.write("| " + " | ".join(headers) + " |\n")
        f.write("| " + " | ".join(["---"] * len(headers)) + " |\n")
        
        for r in rows[1:]:
            vals = [str(c).replace("\n", " <br> ").replace("|", "&#124;") if c is not None else "" for c in r]
            f.write("| " + " | ".join(vals) + " |\n")
        f.write("\n---\n\n")

print(f"Generated {md_path} with all 19 sheets!")
